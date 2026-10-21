# -*- coding: utf-8 -*-
"""就地修改汇报 PPTX 的指定页（不重建整个 deck，避免覆盖手动改动）。

当前用途：
  第 2 页  -> 精简为「一句话定义 + 前向/反向 + 生成」三行，字号放大，30 秒可讲完
  第 4 页  -> 只保留 Goodfellow 分类图 + 一句说明，

用法:  <bundled python> patch_slides.py [--dry-run]
"""
import copy
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "..", "report-deck", "DDPM_Academic_Report_v2.pptx")

# ---- 与 deck 一致的配色（build_deck.py 的 design system）----
NAVY = RGBColor(0x0B, 0x25, 0x45)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GREY = RGBColor(0x5A, 0x66, 0x72)
LIGHT = RGBColor(0xF2, 0xF5, 0xF8)
FONT = "Calibri"

EYEBROW_Y = 0.40      # kicker 行
TITLE_Y = 0.70        # 主标题行
KEEP_TOP = 1.50       # 低于此高度的形状视为「页眉」，一律保留


def is_picture(shape):
    return shape.shape_type is not None and "PICTURE" in str(shape.shape_type)


def wipe(slide, keep_pictures=True, keep_top=KEEP_TOP):
    """删除正文区全部形状；页眉与页脚（低于 7.0in）保留。"""
    removed = []
    for shape in list(slide.shapes):
        y = shape.top / 914400.0
        h = shape.height / 914400.0
        if y + h <= keep_top:                       # 页眉
            continue
        if y >= 7.00:                               # 页脚
            continue
        if is_picture(shape) and keep_pictures:     # 图片单独按需处理
            continue
        removed.append(shape.shape_type)
        shape._element.getparent().remove(shape._element)
    return removed


def drop_pictures_except(slide, keep_indices):
    """按出现顺序删除图片，只保留指定序号的图片。"""
    pics = [sh for sh in slide.shapes if is_picture(sh)]
    keep_ids = {id(pics[i]._element) for i in keep_indices if 0 <= i < len(pics)}
    n = 0
    for p in pics:
        if id(p._element) not in keep_ids:
            p._element.getparent().remove(p._element)
            n += 1
    return n


def textbox(slide, x, y, w, h, text, size, bold=False, color=NAVY,
            line_spacing=1.0, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
    return tb


def set_picture(slide, pic, x, y, w):
    """等比缩放并移动到 (x, y)，宽度为 w。"""
    iw, ih = pic.width, pic.height
    new_h = int(w * 914400.0 * (ih / iw))
    pic.left, pic.top = Inches(x), Inches(y)
    pic.width = Inches(w)
    pic.height = new_h
    return new_h / 914400.0


# ==================================================================
# 第 2 页：精简 + 放大字号（目标：30 秒讲完）
# ==================================================================
def patch_page2(slide):
    wipe(slide, keep_pictures=True)
    drop_pictures_except(slide, keep_indices=[0])       # 只留猫的那张过程图

    # 三行文字，19pt：约 6 行 + 段间距，垂直居中在标题与页脚之间
    textbox(slide, 0.62, 1.95, 6.95, 3.90,
            "Diffusion models are inspired by a familiar physical "
            "process: particles drift from high to low concentration until "
            "the system reaches equilibrium.\n"
            "\n"
            "A diffusion model borrows this idea. It adds noise to an image "
            "step by step, then trains a network to reverse that process.\n"
            "\n"
            "Once the network can turn noise back into an image, we can "
            "generate new samples from pure random noise \u2014 no existing "
            "image required.",
            size=19, color=NAVY, line_spacing=1.22)

    pics = [sh for sh in slide.shapes if is_picture(sh)]
    if pics:
        set_picture(slide, pics[0], 8.10, 2.30, 4.60)


# ==================================================================
# 第 4 页：只留分类图 + 一句说明
# ==================================================================
def patch_page4(slide):
    wipe(slide, keep_pictures=True)
    drop_pictures_except(slide, keep_indices=[0])       # 只留 Goodfellow 分类图

    pics = [sh for sh in slide.shapes if is_picture(sh)]
    if pics:
        # 左对齐放大到 7.00in 宽 -> 高约 3.30in；为下方文字与页脚留出空间
        set_picture(slide, pics[0], 0.62, 1.56, 7.00)
        bottom = 1.56 + pics[0].height / 914400.0
    else:
        bottom = 5.10

    textbox(slide, 0.62, bottom + 0.34, 7.60, 1.10,
            "Diffusion sits under Explicit density \u2192 Approximate "
            "density: a generative model reached by an indirect, iterative "
            "sampling procedure.",
            size=19, color=NAVY, line_spacing=1.24)
    textbox(slide, 8.35, 5.95, 4.35, 1.00,
            "Figure adapted from Goodfellow,\n"
            "Tutorial on Generative Adversarial Networks, 2017",
            size=11.5, color=GREY, line_spacing=1.16)


def main():
    dry = "--dry-run" in sys.argv
    prs = Presentation(DECK)
    s2, s4 = prs.slides[1], prs.slides[3]
    patch_page2(s2)
    patch_page4(s4)
    print(("DRY-RUN " if dry else "") + "已修改第 2、4 页")
    if not dry:
        prs.save(DECK)
        print("已保存:", os.path.abspath(DECK))


if __name__ == "__main__":
    main()
